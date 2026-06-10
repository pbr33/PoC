"""Document Processor — Ingestion & Intelligence Layer.

Extracts text from multiple file formats (PDF, DOCX, XLSX, PPTX, TXT, CSV)
and performs structural analysis for downstream agents.
"""

import io
import re


class DocumentProcessor:
    """Multi-format document text extractor and analyzer."""

    def extract_text(self, uploaded_file) -> str:
        """Extract text from an uploaded Streamlit file object."""
        name = uploaded_file.name.lower()
        raw_bytes = uploaded_file.read()
        uploaded_file.seek(0)

        if name.endswith(".pdf"):
            return self._extract_pdf(raw_bytes)
        elif name.endswith(".docx"):
            return self._extract_docx(raw_bytes)
        elif name.endswith(".xlsx") or name.endswith(".xls"):
            return self._extract_xlsx(raw_bytes)
        elif name.endswith(".pptx"):
            return self._extract_pptx(raw_bytes)
        elif name.endswith(".csv"):
            return self._extract_csv(raw_bytes)
        elif name.endswith(".txt"):
            return raw_bytes.decode("utf-8", errors="replace")
        else:
            return raw_bytes.decode("utf-8", errors="replace")

    def _extract_pdf(self, data: bytes) -> str:
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(io.BytesIO(data))
            text = ""
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n\n"
            return text.strip()
        except Exception as e:
            return f"[PDF extraction error: {e}]"

    def _extract_docx(self, data: bytes) -> str:
        try:
            from docx import Document
            doc = Document(io.BytesIO(data))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            # Also get table content
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        paragraphs.append(" | ".join(cells))
            return "\n".join(paragraphs)
        except Exception as e:
            return f"[DOCX extraction error: {e}]"

    def _extract_xlsx(self, data: bytes) -> str:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            text_parts = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                text_parts.append(f"=== Sheet: {sheet_name} ===")
                for row in ws.iter_rows(max_row=500, values_only=True):
                    cells = [str(c) if c is not None else "" for c in row]
                    if any(cells):
                        text_parts.append(" | ".join(cells))
            return "\n".join(text_parts)
        except Exception as e:
            return f"[XLSX extraction error: {e}]"

    def _extract_pptx(self, data: bytes) -> str:
        try:
            from pptx import Presentation
            prs = Presentation(io.BytesIO(data))
            text_parts = []
            for i, slide in enumerate(prs.slides, 1):
                text_parts.append(f"=== Slide {i} ===")
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text_parts.append(shape.text)
            return "\n".join(text_parts)
        except Exception as e:
            return f"[PPTX extraction error: {e}]"

    def _extract_csv(self, data: bytes) -> str:
        try:
            text = data.decode("utf-8", errors="replace")
            return text[:50000]
        except Exception as e:
            return f"[CSV extraction error: {e}]"

    def analyze_structure(self, text: str) -> dict:
        """Analyze the structural characteristics of the document text."""
        lines = text.split("\n")
        non_empty = [l for l in lines if l.strip()]

        # Detect sections via heading patterns
        sections = []
        for line in non_empty:
            stripped = line.strip()
            # Detect numbered sections, all-caps headings, markdown headings
            if (
                re.match(r"^\d+[\.\)]\s+\w", stripped)
                or (stripped.isupper() and len(stripped) > 3 and len(stripped) < 80)
                or re.match(r"^#{1,4}\s+", stripped)
                or re.match(r"^(Section|Chapter|Part)\s+\d", stripped, re.IGNORECASE)
            ):
                sections.append(stripped[:100])

        # Word / character counts
        words = text.split()
        tables_detected = text.count("|") > 10

        # Detect key entities
        tech_keywords = [
            "azure", "aws", "cloud", "api", "database", "sql", "python", "react",
            "sharepoint", "teams", "power bi", "kubernetes", "docker", ".net",
            "java", "node", "javascript", "typescript", "microservices", "serverless",
            "devops", "ci/cd", "machine learning", "ai", "cosmos", "blob storage",
        ]
        found_tech = [kw for kw in tech_keywords if kw in text.lower()]

        return {
            "total_lines": len(lines),
            "non_empty_lines": len(non_empty),
            "word_count": len(words),
            "character_count": len(text),
            "section_count": len(sections),
            "sections": sections[:30],
            "has_tables": tables_detected,
            "technologies_mentioned": found_tech,
            "confidence_score": min(0.95, 0.5 + len(non_empty) * 0.001),
        }
