# ═══════════════════════════════════════════════════════════════════════
#  AGENT TRAINING — instruction store + feedback loop
#  All data lives in the shared SQLite DB (training_instructions +
#  estimation_feedback tables).  load_training_context() is called
#  at the RAG step of every pipeline run and the returned string is
#  prepended to every agent's system prompt.
# ═══════════════════════════════════════════════════════════════════════
import sqlite3, json, os
from datetime import datetime

from .database import _DB_PATH


# ── Read ──────────────────────────────────────────────────────────────

def load_training_context(project_types: list = None) -> str:
    """
    Load active instructions + recent feedback summaries from SQLite.
    project_types: list of detected types for this run (e.g. ["AI", "Cloud"]).
      Instructions tagged with any of those types are included.
      Instructions tagged [] (general) always load regardless.
    Returns a formatted block ready to prepend to agent system prompts.
    Returns "" if nothing is stored yet.
    """
    try:
        con = sqlite3.connect(_DB_PATH)
        con.row_factory = sqlite3.Row

        all_instrs = con.execute(
            "SELECT instruction, category, project_types FROM training_instructions "
            "WHERE active=1 ORDER BY id ASC"
        ).fetchall()

        # Filter by project type: empty project_types on instruction = applies to all
        instructions = []
        for row in all_instrs:
            try:
                ipt = json.loads(row["project_types"] or "[]")
            except Exception:
                ipt = []
            if not ipt:  # general — always include
                instructions.append(row)
            elif not project_types:  # caller didn't specify type — include everything
                instructions.append(row)
            elif any(pt in ipt for pt in project_types):
                instructions.append(row)

        feedback = con.execute(
            "SELECT client_name, bella_hours, actual_hours, bella_cost, actual_cost, "
            "phase_deltas, notes FROM estimation_feedback "
            "WHERE active=1 ORDER BY id DESC LIMIT 20"
        ).fetchall()
        con.close()
    except Exception:
        return ""

    if not instructions and not feedback:
        return ""

    parts = ["═" * 60,
             "AGENT TRAINING CONTEXT — read and apply before estimating",
             "═" * 60]

    if instructions:
        parts.append("\n📋 INSTRUCTIONS FROM PRESALES LEAD:")
        by_cat: dict = {}
        for row in instructions:
            cat = (row["category"] or "general").title()
            by_cat.setdefault(cat, []).append(row["instruction"])
        for cat, items in by_cat.items():
            parts.append(f"\n[{cat}]")
            for item in items:
                parts.append(f"  • {item}")

    if feedback:
        parts.append("\n\n📊 PAST ESTIMATE vs ACTUAL (learn from these gaps):")
        for fb in feedback:
            h_delta = fb["actual_hours"] - fb["bella_hours"]
            c_delta = fb["actual_cost"]  - fb["bella_cost"]
            direction_h = "over-estimated" if h_delta < 0 else "under-estimated"
            direction_c = "over-estimated" if c_delta < 0 else "under-estimated"
            client = fb["client_name"] or "Client"
            parts.append(
                f"\n  Project: {client}\n"
                f"    Hours  — BELLA: {fb['bella_hours']}h | Actual: {fb['actual_hours']}h "
                f"({direction_h} by {abs(h_delta)}h)\n"
                f"    Cost   — BELLA: ${fb['bella_cost']:,}/mo | Actual: ${fb['actual_cost']:,}/mo "
                f"({direction_c} by ${abs(c_delta):,}/mo)"
            )
            try:
                phases = json.loads(fb["phase_deltas"] or "[]")
                if phases:
                    parts.append("    Phase breakdown:")
                    for p in phases[:8]:
                        d = p.get("actual_hours", 0) - p.get("bella_hours", 0)
                        sign = "+" if d >= 0 else ""
                        parts.append(f"      {p.get('phase','?')}: BELLA {p.get('bella_hours',0)}h → Actual {p.get('actual_hours',0)}h ({sign}{d}h)")
            except Exception:
                pass
            if fb["notes"]:
                parts.append(f"    Note: {fb['notes']}")

        # Compute aggregate bias
        total_h_delta = sum(fb["actual_hours"] - fb["bella_hours"] for fb in feedback)
        avg_h_delta   = round(total_h_delta / len(feedback))
        if abs(avg_h_delta) > 10:
            bias = "tends to UNDER-estimate" if avg_h_delta > 0 else "tends to OVER-estimate"
            parts.append(
                f"\n  ⚠️  SYSTEMATIC BIAS: Based on {len(feedback)} past project(s), BELLA {bias} "
                f"hours by an average of {abs(avg_h_delta)}h. Adjust your estimates accordingly."
            )

    parts.append("\n" + "═" * 60)
    return "\n".join(parts)


# ── Write ─────────────────────────────────────────────────────────────

def add_instruction(instruction: str, category: str = "general",
                    created_by: str = "", project_types: list = None) -> int:
    """Insert a new training instruction. Returns the new row id."""
    con = sqlite3.connect(_DB_PATH)
    cur = con.execute(
        "INSERT INTO training_instructions (created_at, instruction, category, created_by, project_types) "
        "VALUES (?,?,?,?,?)",
        (datetime.utcnow().isoformat(), instruction.strip(), category.lower(), created_by,
         json.dumps(project_types or []))
    )
    con.commit()
    row_id = cur.lastrowid
    con.close()
    return row_id


def add_feedback(client_name: str, bella_hours: int, actual_hours: int,
                 bella_cost: int, actual_cost: int,
                 phase_deltas: list, notes: str = "",
                 proposal_id: str = "", created_by: str = "") -> int:
    """Store one feedback record comparing BELLA estimate vs actual Excel."""
    con = sqlite3.connect(_DB_PATH)
    cur = con.execute(
        "INSERT INTO estimation_feedback "
        "(created_at, proposal_id, client_name, bella_hours, actual_hours, "
        " bella_cost, actual_cost, phase_deltas, notes, created_by) "
        "VALUES (?,?,?,?,?,?,?,?,?,?)",
        (datetime.utcnow().isoformat(), proposal_id, client_name,
         int(bella_hours), int(actual_hours), int(bella_cost), int(actual_cost),
         json.dumps(phase_deltas), notes, created_by)
    )
    con.commit()
    row_id = cur.lastrowid
    con.close()
    return row_id


def list_instructions(active_only: bool = False) -> list:
    con = sqlite3.connect(_DB_PATH)
    con.row_factory = sqlite3.Row
    q = "SELECT * FROM training_instructions"
    if active_only:
        q += " WHERE active=1"
    q += " ORDER BY id DESC"
    rows = [dict(r) for r in con.execute(q).fetchall()]
    con.close()
    return rows


def list_feedback() -> list:
    con = sqlite3.connect(_DB_PATH)
    con.row_factory = sqlite3.Row
    rows = [dict(r) for r in con.execute(
        "SELECT * FROM estimation_feedback ORDER BY id DESC"
    ).fetchall()]
    con.close()
    return rows


def toggle_instruction(row_id: int, active: bool):
    con = sqlite3.connect(_DB_PATH)
    con.execute("UPDATE training_instructions SET active=? WHERE id=?", (int(active), row_id))
    con.commit()
    con.close()


def delete_instruction(row_id: int):
    con = sqlite3.connect(_DB_PATH)
    con.execute("DELETE FROM training_instructions WHERE id=?", (row_id,))
    con.commit()
    con.close()


def delete_feedback(row_id: int):
    con = sqlite3.connect(_DB_PATH)
    con.execute("DELETE FROM estimation_feedback WHERE id=?", (row_id,))
    con.commit()
    con.close()


# ── Excel parser ──────────────────────────────────────────────────────

def parse_excel_estimate(file_bytes: bytes) -> dict:
    """
    Parse an Excel file (the actual estimate sent to client).
    Returns:
      {
        "total_hours": int,
        "total_cost": int,
        "phases": [{"phase": str, "actual_hours": int, "actual_cost": int}]
      }
    Handles common layouts: one column for phase/task, one for hours, one for cost.
    """
    import io
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
    except ImportError:
        # Fall back to xlrd / pandas if openpyxl not available
        try:
            import pandas as pd
            df = pd.read_excel(io.BytesIO(file_bytes))
            return _parse_df(df)
        except Exception:
            return {"total_hours": 0, "total_cost": 0, "phases": []}
    except Exception:
        return {"total_hours": 0, "total_cost": 0, "phases": []}

    # Find header row — look for cells containing "hour", "day", "effort", "cost"
    header_row = None
    hour_col = cost_col = phase_col = None
    for row_idx, row in enumerate(ws.iter_rows(max_row=20), 1):
        for cell in row:
            val = str(cell.value or "").lower().strip()
            if any(k in val for k in ["hour", "days", "effort", "man day"]):
                hour_col = cell.column
                header_row = row_idx
            if any(k in val for k in ["cost", "price", "amount", "budget"]):
                cost_col = cell.column
                header_row = row_idx
            if any(k in val for k in ["phase", "activity", "task", "module", "workstream"]):
                phase_col = cell.column
                header_row = row_idx

    phases = []
    total_hours = 0
    total_cost  = 0

    if header_row and (hour_col or cost_col):
        for row in ws.iter_rows(min_row=header_row + 1):
            phase_name = str(row[phase_col - 1].value or "").strip() if phase_col else f"Row {row[0].row}"
            hours_val  = row[hour_col - 1].value if hour_col else 0
            cost_val   = row[cost_col  - 1].value if cost_col  else 0
            try:
                h = float(hours_val or 0)
                c = float(cost_val  or 0)
            except (ValueError, TypeError):
                continue
            if h == 0 and c == 0:
                continue
            # If "days" column, convert to hours (×8)
            if hour_col:
                hdr_val = str(ws.cell(header_row, hour_col).value or "").lower()
                if "day" in hdr_val:
                    h = h * 8
            phases.append({"phase": phase_name, "actual_hours": int(h), "actual_cost": int(c)})
            total_hours += int(h)
            total_cost  += int(c)
    else:
        # Fallback: sum all numeric cells
        for row in ws.iter_rows():
            for cell in row:
                try:
                    if isinstance(cell.value, (int, float)) and cell.value > 0:
                        total_hours += int(cell.value)
                except Exception:
                    pass

    return {"total_hours": total_hours, "total_cost": total_cost, "phases": phases}


def find_similar_instructions(new_text: str, project_types: list = None) -> list:
    """
    Fast text-overlap check — no AI call, no latency.
    Returns list of existing instructions with word-overlap > 50% with new_text.
    Each entry: {"id": int, "instruction": str, "overlap": float}
    """
    new_words = set(new_text.lower().split())
    if len(new_words) < 4:
        return []
    con = sqlite3.connect(_DB_PATH)
    con.row_factory = sqlite3.Row
    rows = [dict(r) for r in con.execute(
        "SELECT id, instruction, project_types FROM training_instructions WHERE active=1"
    ).fetchall()]
    con.close()
    matches = []
    for row in rows:
        ex_words = set(row["instruction"].lower().split())
        if not ex_words:
            continue
        overlap = len(new_words & ex_words) / max(len(new_words | ex_words), 1)
        if overlap > 0.50:
            matches.append({"id": row["id"], "instruction": row["instruction"], "overlap": round(overlap, 2)})
    return sorted(matches, key=lambda x: x["overlap"], reverse=True)


def consolidate_instructions(project_types: list = None) -> dict:
    """
    AI pass: group similar/overlapping instructions and replace each group
    with one canonical merged rule. Returns {"before": int, "after": int, "groups_merged": int}.
    Skips if fewer than 6 instructions (nothing useful to consolidate).
    """
    con = sqlite3.connect(_DB_PATH)
    con.row_factory = sqlite3.Row
    rows = [dict(r) for r in con.execute(
        "SELECT id, instruction, category, project_types FROM training_instructions WHERE active=1 ORDER BY id ASC"
    ).fetchall()]
    con.close()

    if len(rows) < 6:
        return {"before": len(rows), "after": len(rows), "groups_merged": 0, "skipped": True}

    try:
        from .ai_clients import AnthropicAI
        client = AnthropicAI.from_session()
        if not client.is_live:
            return {"before": len(rows), "after": len(rows), "groups_merged": 0, "error": "AI offline"}
    except Exception as e:
        return {"before": len(rows), "after": len(rows), "groups_merged": 0, "error": str(e)}

    instrs_text = "\n".join(
        f"[ID:{r['id']}][{r['category']}] {r['instruction']}" for r in rows
    )
    system = (
        "You consolidate training instructions for an AI estimation tool. "
        "Find groups of instructions that are redundant, overlapping, or say the same thing differently. "
        "For each group, write ONE canonical instruction that captures the intent of all members. "
        "DO NOT group instructions about different topics just because they are related. "
        "Only group if two+ instructions are genuinely redundant (merging them loses no information). "
        "Singleton instructions (no overlap with others) must NOT appear in any group. "
        "Return ONLY valid JSON (no markdown):\n"
        '{"groups": [{"canonical": "...", "category": "general|hours|streams|tasks", '
        '"ids_to_replace": [int, ...]}]}'
    )
    result = client._call(system, instrs_text, max_tokens=1200)

    if not (result and isinstance(result, dict) and "groups" in result):
        return {"before": len(rows), "after": len(rows), "groups_merged": 0, "error": "AI returned no groups"}

    groups = [g for g in result.get("groups", []) if len(g.get("ids_to_replace", [])) >= 2]
    if not groups:
        return {"before": len(rows), "after": len(rows), "groups_merged": 0}

    ids_by_row = {r["id"]: r for r in rows}
    ids_to_disable = []
    new_entries = []
    for g in groups:
        canonical = (g.get("canonical") or "").strip()
        cat = g.get("category", "general")
        ids = [i for i in g.get("ids_to_replace", []) if i in ids_by_row]
        if not canonical or len(ids) < 2:
            continue
        # Inherit project_types from the first matched instruction
        pts = []
        try:
            pts = json.loads(ids_by_row[ids[0]].get("project_types") or "[]")
        except Exception:
            pass
        ids_to_disable.extend(ids)
        new_entries.append((canonical, cat, pts))

    if not ids_to_disable:
        return {"before": len(rows), "after": len(rows), "groups_merged": 0}

    con = sqlite3.connect(_DB_PATH)
    for id_ in set(ids_to_disable):
        con.execute("UPDATE training_instructions SET active=0 WHERE id=?", (id_,))
    for canonical, cat, pts in new_entries:
        con.execute(
            "INSERT INTO training_instructions (created_at, instruction, category, created_by, project_types) "
            "VALUES (?,?,?,?,?)",
            (datetime.utcnow().isoformat(), canonical, cat, "system-consolidation", json.dumps(pts))
        )
    con.commit()
    con.close()

    after = len(rows) - len(set(ids_to_disable)) + len(new_entries)
    return {"before": len(rows), "after": after, "groups_merged": len(new_entries)}


def extract_structural_patterns(file_bytes: bytes) -> dict:
    """
    Read an Excel estimation file and extract STRUCTURAL patterns (not hours).
    Returns a dict with sheet names, roles, technologies, and per-stream task samples.
    Feed the result into generate_instructions_from_excel().
    """
    import io, re as _re
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception:
        return {}

    sheet_names = wb.sheetnames

    # Estimation sheets: any sheet whose name contains these substrings
    estimation_sheets = [
        s for s in sheet_names
        if any(k in s.lower() for k in ["estimat", "effort", "dev", "qa", "infra", "visual", "platform"])
    ]

    _TECH_KW = [
        "power bi", "power automate", "power apps", "power platform",
        "microsoft fabric", "fabric", "lakehouse", "dataflow", "data factory", "adf",
        "azure sql", "azure blob", "azure vm", "azure data", "azure",
        "sql server", "synapse", "databricks", "spark",
        "sharepoint", "teams", "copilot", "viva",
        "react", "angular", "vue", "node.js", ".net", "python",
        "tableau", "qlik", "looker",
        "rpa", "automate desktop", "uipath", "blue prism",
        "entra id", "entra", "active directory", "key vault", "devops",
        "kubernetes", "docker", "container",
        "snowflake", "dbt", "redshift", "bigquery",
        "openai", "llm", "machine learning",
        "rest api", "graphql", "api integration",
    ]

    tech_found: set = set()
    streams: dict = {}

    for sname in estimation_sheets:
        ws = wb[sname]
        tasks = []
        for row in ws.iter_rows(min_row=1, max_row=250):
            cell_val = str(row[0].value or "").strip()
            # Skip empty, numeric-only, or header-like cells
            if not cell_val or _re.match(r'^[\d.,\s]+$', cell_val) or len(cell_val) <= 2:
                continue
            if cell_val.lower() in ("task", "phase", "activity", "module", "#", "no", "no.", "sr"):
                continue
            tasks.append(cell_val)
            # Scan full row for tech keywords
            row_text = " ".join(str(c.value or "").lower() for c in row)
            for kw in _TECH_KW:
                if kw in row_text:
                    tech_found.add(kw)
        if tasks:
            streams[sname] = tasks[:25]

    # Scan ALL sheets for tech keywords (architecture, flow diagrams have text too)
    for sname in sheet_names:
        if sname in estimation_sheets:
            continue
        ws = wb[sname]
        for row in ws.iter_rows(max_row=100):
            row_text = " ".join(str(c.value or "").lower() for c in row)
            for kw in _TECH_KW:
                if kw in row_text:
                    tech_found.add(kw)

    # Extract roles from Summary / Team sheet
    team_roles: list = []
    _role_kw = ["engineer", "developer", "architect", "manager", "lead", "analyst", "qa", "tester", "consultant"]
    for sname in sheet_names:
        if not any(k in sname.lower() for k in ["summary", "team", "resource", "allocation"]):
            continue
        ws = wb[sname]
        for row in ws.iter_rows(max_row=60):
            first = str(row[0].value or "").strip()
            if not first or len(first) < 4:
                continue
            if any(k in first.lower() for k in _role_kw):
                if first not in team_roles:
                    team_roles.append(first)

    # Fallback: infer roles from estimation sheet names
    inferred_roles = []
    for sname in estimation_sheets:
        cleaned = _re.sub(r'(?i)estimation|estimate|effort', '', sname).strip(" -_")
        if cleaned and cleaned not in inferred_roles:
            inferred_roles.append(cleaned)

    roles = list(dict.fromkeys(team_roles + inferred_roles))

    return {
        "sheet_names": sheet_names,
        "estimation_sheets": estimation_sheets,
        "roles": roles,
        "streams": streams,
        "technologies": sorted(tech_found),
    }


def generate_instructions_from_excel(patterns: dict, project_context: str = "") -> list:
    """
    Call Claude to turn structural Excel patterns into training instructions for BELLA.
    Returns [{"text": str, "category": str}].
    """
    if not patterns:
        return []
    try:
        from .ai_clients import AnthropicAI
        client = AnthropicAI.from_session()
        if not client.is_live:
            return []
    except Exception:
        return []

    lines = []
    if patterns.get("sheet_names"):
        lines.append(f"Excel sheets: {', '.join(patterns['sheet_names'])}")
    if patterns.get("estimation_sheets"):
        lines.append(f"Role/estimation sheets (one sheet per role/team): {', '.join(patterns['estimation_sheets'])}")
    if patterns.get("roles"):
        lines.append(f"Team roles identified: {', '.join(patterns['roles'][:15])}")
    if patterns.get("technologies"):
        lines.append(f"Technologies mentioned in the file: {', '.join(patterns['technologies'][:25])}")
    for sname, tasks in (patterns.get("streams") or {}).items():
        lines.append(f"\n[{sname}] phases/tasks:\n  " + "\n  ".join(tasks[:10]))
    if project_context:
        lines.append(f"\nProject context: {project_context}")

    system = (
        "You are a presales estimation expert analyzing a real client Excel file. "
        "Extract TRAINING INSTRUCTIONS for an AI estimation tool called BELLA. "
        "Instructions must describe STRUCTURE, ROLES, STREAMS, and METHODOLOGY — NOT specific hours or costs.\n\n"
        "GOOD instruction examples:\n"
        "• 'For Power BI projects, always include a Visualization Developer stream separate from Data Engineering'\n"
        "• 'For Power Automate Desktop/RPA projects, include a Power Platform Developer stream with bot dev phases'\n"
        "• 'Use Bronze/Silver/Gold medallion architecture phases in Data Engineering stream for Fabric/Lakehouse projects'\n"
        "• 'ETL QA is always a separate stream — never bundle it with the development stream'\n"
        "• 'For projects with 4+ data sources, list a per-source ingestion task in the Data Engineering stream'\n"
        "• 'Always include a Wireframes & UX phase before dashboard development in BI projects'\n\n"
        "BAD instructions (do NOT generate):\n"
        "• 'Power BI development takes 70 hours' — this is an hours estimate, not a structural rule\n"
        "• 'Allocate 50% to Data Engineer' — too specific to one project\n\n"
        "Return ONLY valid JSON (no markdown fence, no extra text):\n"
        '{"instructions": [{"text": "...", "category": "streams|tasks|roles|general"}]}'
    )

    result = client._call(system, "\n".join(lines), max_tokens=1800)

    if not (result and isinstance(result, dict) and "instructions" in result):
        return []

    return [
        {"text": i.get("text", "").strip(), "category": i.get("category", "general")}
        for i in result.get("instructions", [])
        if i.get("text", "").strip()
    ]


def _parse_df(df) -> dict:
    total_hours = 0
    total_cost  = 0
    phases = []
    hour_col = next((c for c in df.columns if any(k in str(c).lower() for k in ["hour","effort","day"])), None)
    cost_col = next((c for c in df.columns if any(k in str(c).lower() for k in ["cost","price","amount"])), None)
    name_col = next((c for c in df.columns if any(k in str(c).lower() for k in ["phase","task","activity","module"])), None)
    for _, row in df.iterrows():
        h = int(float(row.get(hour_col, 0) or 0)) if hour_col else 0
        c = int(float(row.get(cost_col, 0)  or 0)) if cost_col  else 0
        n = str(row.get(name_col, "")) if name_col else ""
        if h == 0 and c == 0:
            continue
        phases.append({"phase": n, "actual_hours": h, "actual_cost": c})
        total_hours += h
        total_cost  += c
    return {"total_hours": total_hours, "total_cost": total_cost, "phases": phases}
