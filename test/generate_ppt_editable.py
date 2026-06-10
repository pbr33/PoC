"""
GLO Finance AS-IS Process Map — fully editable PPTX
All elements are native PowerPoint shapes (rectangles, diamonds, connectors, text).
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from lxml import etree

OUTPUT = Path(r"c:\Users\prabhakargupta\Downloads\dream\test\GLO_Finance_ASIS_Editable.pptx")

# ── Diagram constants (match v3 HTML) ────────────────────────────────────────
CANVAS_W = 3540
CANVAS_H = 1300
LABEL_W  = 210
BW, BH   = 148, 46   # box width / height
DW       = 68         # diamond half-width (from centre to vertex)
EW, EH   = 148, 36   # exception box

# ── Colours ──────────────────────────────────────────────────────────────────
def rgb(h):
    h = h.lstrip('#')
    return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))

C_BG    = rgb('#03060F')
C_PANEL = rgb('#0D1117')
C_WHITE = rgb('#FFFFFF')
C_GREY  = rgb('#94A3B8')
C_TEAL  = rgb('#00E5C0')
C_AMBER = rgb('#F59E0B')
C_ORANGE= rgb('#F97316')
C_GREEN = rgb('#4ade80')

CONN_COLORS = {
    'main': rgb('#94A3B8'),
    'yes' : C_GREEN,
    'no'  : C_ORANGE,
    'loop': C_ORANGE,
    'hand': rgb('#64748B'),
}

# ── Lane definitions ─────────────────────────────────────────────────────────
LANES = [
    {'id':'mdm',     'label':'Supplier &\nFactory MDM',     'color':'#00E5C0',
     'y_top':76,  'y_bot':240,  'systems':['EQOS','NeoGrid','SSUP Portal']},
    {'id':'po',      'label':'Purchase Order\nCreation',    'color':'#4FB3FF',
     'y_top':240, 'y_bot':400,  'systems':['LM / MM','ISOS','OTM / BizTalk']},
    {'id':'ship',    'label':'Shipment\nExecution',         'color':'#A78BFA',
     'y_top':400, 'y_bot':550,  'systems':['OTM','BizTalk / RTA','FF Systems']},
    {'id':'prodinv', 'label':'Product Invoice\nProcessing', 'color':'#FFD04A',
     'y_top':550, 'y_bot':775,  'systems':['ISOS / OTM BM','Documentum','EDI → OFI','MS Access']},
    {'id':'freight', 'label':'Freight\nSelf-Billing',       'color':'#FF6B6B',
     'y_top':775, 'y_bot':970,  'systems':['MS Access DB','Rolling File','Oracle Fusion','Tungsten']},
    {'id':'ar',      'label':'AR & Intercompany\nBilling',  'color':'#C57BFF',
     'y_top':970, 'y_bot':1110, 'systems':['Oracle Fusion AR','RIMS (CE)','Oracle R12']},
    {'id':'customs', 'label':'Customs &\nDuty',             'color':'#52E67A',
     'y_top':1110,'y_bot':1295, 'systems':['OTM CSS/IMS','NBS/CM-UK','WTA Portal','DDA (TISL)']},
]

PHASES = [
    {'label':'Supplier\nOnboarding',         'x_start':210,  'x_end':590 },
    {'label':'Purchase Order\nCreation',     'x_start':590,  'x_end':960 },
    {'label':'Shipment\nExecution',          'x_start':960,  'x_end':1310},
    {'label':'Invoice\nPreparation',         'x_start':1310, 'x_end':1680},
    {'label':'Cost Validation\n& Agreement', 'x_start':1680, 'x_end':2060},
    {'label':'PO Raise &\nSubmission',       'x_start':2060, 'x_end':2460},
    {'label':'Payment &\nAR Billing',        'x_start':2460, 'x_end':2850},
    {'label':'Customs Clearance\n& Duty',    'x_start':2850, 'x_end':3520},
]

ELEMS = [
    # MDM lane
    {'id':'mdm_ssup','lane':'mdm','cx':340, 'type':'box','bold':True,
     'label':'SSUP: Collect Docs\n& Submit via NeoGrid','sys':'NeoGrid · SSUP Portal'},
    {'id':'mdm_val', 'lane':'mdm','cx':520, 'type':'dia',
     'label':'MDM Validation\nOK?','sys':'EQOS · NeoGrid'},
    {'id':'mdm_gap', 'lane':'mdm','cx':520, 'cy_off':62,'type':'exc',
     'label':'Gaps Flagged\n↑ SSUP Resolves','sys':'NeoGrid'},
    {'id':'mdm_eqos','lane':'mdm','cx':710, 'type':'out','bold':True,
     'label':'Supplier/Factory\nCreated in EQOS','sys':'EQOS'},
    {'id':'mdm_feed','lane':'mdm','cx':890, 'type':'box',
     'label':'EQOS → OFI (weekly)\n+ ISOS (15-min push)','sys':'EQOS → OFI · ISOS'},
    # PO lane
    {'id':'po_lm',    'lane':'po','cx':700,  'type':'box','bold':True,
     'label':'Buyer Creates PO\nin LM / MM','sys':'LM · MM System'},
    {'id':'po_isos',  'lane':'po','cx':900,  'type':'dia',
     'label':'ISOS: Factory\nCompliant & Valid?','sys':'ISOS · EQOS'},
    {'id':'po_cancel','lane':'po','cx':900,  'cy_off':60,'type':'exc',
     'label':'ISOS Cancels PO\n(LM not notified)','sys':'ISOS'},
    {'id':'po_enrich','lane':'po','cx':1090, 'type':'box',
     'label':'ISOS Enriches PO\n+ Assigns FF/LSP','sys':'ISOS'},
    {'id':'po_otm',   'lane':'po','cx':1270, 'type':'out','bold':True,
     'label':'PO → OTM +\nSupplier (PDF via OTM)','sys':'OTM · BizTalk'},
    # Shipment lane
    {'id':'sh_fob', 'lane':'ship','cx':1060,'type':'box',
     'label':'Origin FF: FOB\nHandover → SM1','sys':'FF Systems · BizTalk'},
    {'id':'sh_isos','lane':'ship','cx':1230,'type':'box',
     'label':'ISOS: SM1 Processed\n(version filtering)','sys':'ISOS · BizTalk'},
    {'id':'sh_sm2', 'lane':'ship','cx':1400,'type':'out','bold':True,
     'label':'SM2 → OTM:\nShipment Confirmed','sys':'ISOS → OTM'},
    {'id':'sh_rpt', 'lane':'ship','cx':1580,'type':'box',
     'label':'OTM Daily Inbound\nAll File Generated','sys':'OTM'},
    # Product Invoice lane
    {'id':'inv_iv0', 'lane':'prodinv','cx':1430,'type':'box',
     'label':'FF → IV0 via\nBizTalk → ISOS','sys':'BizTalk → ISOS'},
    {'id':'inv_doc', 'lane':'prodinv','cx':1620,'type':'box',
     'label':'Supplier PDF →\nDocumentum (manual)','sys':'Documentum'},
    {'id':'inv_cost','lane':'prodinv','cx':1800,'type':'dia',
     'label':'ISOS: PO Cost\n= IV0 Cost?','sys':'ISOS · OTM BM'},
    {'id':'inv_bm',  'lane':'prodinv','cx':1800,'cy_off':82,'type':'exc',
     'label':'OTM BM: Buyer\nResolves Cost','sys':'OTM · BizTalk'},
    {'id':'inv_78',  'lane':'prodinv','cx':1990,'type':'box',
     'label':'ISOS 78-col Extract\n(Approved POs only)','sys':'ISOS'},
    {'id':'inv_edi', 'lane':'prodinv','cx':2170,'type':'out',
     'label':'EDI: 78→69-col\nCSV (OFI format)','sys':'EDI'},
    {'id':'inv_3way','lane':'prodinv','cx':2360,'type':'dia',
     'label':'MS Access 3-way\nMatch OK?','sys':'MS Access · Documentum'},
    {'id':'inv_dq',  'lane':'prodinv','cx':2360,'cy_off':82,'type':'exc',
     'label':'DQ Exception:\nFF / Supplier Corrects','sys':'DQ Process'},
    {'id':'inv_ofi', 'lane':'prodinv','cx':2540,'type':'out','bold':True,
     'label':'69-col → OFI:\nAP Invoice Created','sys':'Oracle Financials 11i'},
    {'id':'inv_pay', 'lane':'prodinv','cx':2720,'type':'out','bold':True,
     'label':'HSBC Payment\n(90-day / LC / VF)','sys':'OFI · HSBC Banking'},
    # Freight Self-Billing lane
    {'id':'fr_file', 'lane':'freight','cx':1680,'type':'box',
     'label':'OTM Inbound File\n→ MS Access DB','sys':'OTM → MS Access DB'},
    {'id':'fr_ext',  'lane':'freight','cx':1860,'type':'box',
     'label':'Monday: Extract\nWeekly File / Carrier','sys':'MS Access DB'},
    {'id':'fr_roll', 'lane':'freight','cx':2030,'type':'box',
     'label':'Clean & Load into\nRolling File (Excel)','sys':'Excel Rolling File'},
    {'id':'fr_send', 'lane':'freight','cx':2200,'type':'box',
     'label':'Email Invoice File\nto Carrier for Review','sys':'Outlook · Excel'},
    {'id':'fr_agree','lane':'freight','cx':2380,'type':'dia',
     'label':'Carrier Confirms\nCosts?','sys':'Excel · Email'},
    {'id':'fr_disc', 'lane':'freight','cx':2380,'cy_off':78,'type':'exc',
     'label':'Rate Card Check:\n↑ Discrepancy Resolved','sys':'Rate Card · Email'},
    {'id':'fr_po',   'lane':'freight','cx':2560,'type':'out','bold':True,
     'label':'Oracle Fusion: PO\nRaised (P1002-XXXXXX)','sys':'Oracle Fusion Procurement'},
    {'id':'fr_tung', 'lane':'freight','cx':2730,'type':'box',
     'label':'Carrier Creates\nInvoice via Tungsten','sys':'Tungsten Network'},
    {'id':'fr_pay',  'lane':'freight','cx':2900,'type':'out','bold':True,
     'label':'Tungsten→Fusion AP:\nAP Team Pays (30-60d)','sys':'Oracle Fusion Payables'},
    # AR lane
    {'id':'ar_auto','lane':'ar','cx':2610,'type':'box',
     'label':'OFI: AR Invoice\nAuto-Created (TISL)','sys':'Oracle Financials 11i'},
    {'id':'ar_comm','lane':'ar','cx':2790,'type':'out',
     'label':'Commission Invoice\n(1.5 / 3.5 / 4.25%)','sys':'OFI · TISL'},
    {'id':'ar_dest','lane':'ar','cx':2970,'type':'dia',
     'label':'UK / ROI\nor CE (TICB)?','sys':'OFI · RIMS'},
    {'id':'ar_uk',  'lane':'ar','cx':3150,'cy_off':-36,'type':'box',
     'label':'Oracle Fusion AR\n(UK/ROI) → EDI','sys':'Oracle Fusion AR · EDI'},
    {'id':'ar_ce',  'lane':'ar','cx':3150,'cy_off': 36,'type':'box',
     'label':'CE: RIMS Manual\n→ Oracle R12','sys':'RIMS · Oracle R12'},
    # Customs lane
    {'id':'cus_css',    'lane':'customs','cx':2300,'type':'box',
     'label':'OTM → CSS Auto-sent\nto WTA (pre-arrival)','sys':'OTM → WTA Portal'},
    {'id':'cus_ims',    'lane':'customs','cx':2490,'type':'box',
     'label':'IMS Data: OFI AP\n→ Customs Team','sys':'OFI / ISOS → Customs'},
    {'id':'cus_data',   'lane':'customs','cx':2680,'type':'dia',
     'label':'CSS + IMS\nBoth Available?','sys':'WTA Clearing Agent'},
    {'id':'cus_hold',   'lane':'customs','cx':2680,'cy_off':68,'type':'exc',
     'label':'Container Hold Risk:\n↑ Data Escalated','sys':'WTA · AP Team'},
    {'id':'cus_sfd',    'lane':'customs','cx':2880,'type':'out',
     'label':'WTA Files SFD\nat UK Border','sys':'WTA · HMRC (CHIEF/CDS)'},
    {'id':'cus_nbs',    'lane':'customs','cx':3060,'type':'box',
     'label':'OTM HS Codes\n→ NBS Duty Rates','sys':'OTM → NBS'},
    {'id':'cus_accrual','lane':'customs','cx':3245,'type':'box',
     'label':'Accrual: Duty Rate\nx Value x Qty → GL','sys':'QL Accrual · OFI GL'},
    {'id':'cus_dda',    'lane':'customs','cx':3430,'type':'out','bold':True,
     'label':'TISL Pays Duty\nvia DDA (16th mth)','sys':'DDA · TISL'},
]

CONNS = [
    {'f':'mdm_ssup','fp':'r','t':'mdm_val', 'tp':'l','type':'main','label':''},
    {'f':'mdm_val', 'fp':'r','t':'mdm_eqos','tp':'l','type':'yes', 'label':'YES'},
    {'f':'mdm_val', 'fp':'b','t':'mdm_gap', 'tp':'t','type':'no',  'label':'NO'},
    {'f':'mdm_gap', 'fp':'l','t':'mdm_val', 'tp':'l','type':'loop','label':'retry'},
    {'f':'mdm_eqos','fp':'r','t':'mdm_feed','tp':'l','type':'main','label':''},
    {'f':'mdm_feed','fp':'b','t':'po_isos', 'tp':'t','type':'hand','label':'EQOS→ISOS'},
    {'f':'po_lm',   'fp':'r','t':'po_isos',  'tp':'l','type':'main','label':''},
    {'f':'po_isos', 'fp':'r','t':'po_enrich','tp':'l','type':'yes', 'label':'YES'},
    {'f':'po_isos', 'fp':'b','t':'po_cancel','tp':'t','type':'no',  'label':'NO'},
    {'f':'po_enrich','fp':'r','t':'po_otm',  'tp':'l','type':'main','label':''},
    {'f':'po_otm',  'fp':'b','t':'sh_fob',  'tp':'t','type':'hand','label':'PO→Ship'},
    {'f':'sh_fob',  'fp':'r','t':'sh_isos','tp':'l','type':'main','label':''},
    {'f':'sh_isos', 'fp':'r','t':'sh_sm2', 'tp':'l','type':'main','label':''},
    {'f':'sh_sm2',  'fp':'r','t':'sh_rpt', 'tp':'l','type':'main','label':''},
    {'f':'sh_sm2',  'fp':'b','t':'inv_iv0','tp':'t','type':'hand','label':'IV0 via BizTalk'},
    {'f':'sh_rpt',  'fp':'b','t':'fr_file','tp':'t','type':'hand','label':'OTM All File'},
    {'f':'inv_iv0', 'fp':'r','t':'inv_doc', 'tp':'l','type':'main','label':''},
    {'f':'inv_doc', 'fp':'r','t':'inv_cost','tp':'l','type':'main','label':''},
    {'f':'inv_cost','fp':'r','t':'inv_78',  'tp':'l','type':'yes', 'label':'YES'},
    {'f':'inv_cost','fp':'b','t':'inv_bm',  'tp':'t','type':'no',  'label':'NO'},
    {'f':'inv_bm',  'fp':'l','t':'inv_cost','tp':'l','type':'loop','label':'resolved'},
    {'f':'inv_78',  'fp':'r','t':'inv_edi', 'tp':'l','type':'main','label':''},
    {'f':'inv_edi', 'fp':'r','t':'inv_3way','tp':'l','type':'main','label':''},
    {'f':'inv_3way','fp':'r','t':'inv_ofi', 'tp':'l','type':'yes', 'label':'YES'},
    {'f':'inv_3way','fp':'b','t':'inv_dq',  'tp':'t','type':'no',  'label':'NO'},
    {'f':'inv_dq',  'fp':'l','t':'inv_3way','tp':'l','type':'loop','label':'resubmit'},
    {'f':'inv_ofi', 'fp':'r','t':'inv_pay', 'tp':'l','type':'main','label':''},
    {'f':'inv_ofi', 'fp':'b','t':'ar_auto', 'tp':'t','type':'hand','label':'AP→AR'},
    {'f':'inv_ofi', 'fp':'b','t':'cus_ims', 'tp':'t','type':'hand','label':'IMS→Customs'},
    {'f':'fr_file', 'fp':'r','t':'fr_ext',  'tp':'l','type':'main','label':''},
    {'f':'fr_ext',  'fp':'r','t':'fr_roll', 'tp':'l','type':'main','label':''},
    {'f':'fr_roll', 'fp':'r','t':'fr_send', 'tp':'l','type':'main','label':''},
    {'f':'fr_send', 'fp':'r','t':'fr_agree','tp':'l','type':'main','label':''},
    {'f':'fr_agree','fp':'r','t':'fr_po',   'tp':'l','type':'yes', 'label':'YES'},
    {'f':'fr_agree','fp':'b','t':'fr_disc', 'tp':'t','type':'no',  'label':'NO'},
    {'f':'fr_disc', 'fp':'l','t':'fr_agree','tp':'l','type':'loop','label':'resolved'},
    {'f':'fr_po',   'fp':'r','t':'fr_tung', 'tp':'l','type':'main','label':''},
    {'f':'fr_tung', 'fp':'r','t':'fr_pay',  'tp':'l','type':'main','label':''},
    {'f':'ar_auto', 'fp':'r','t':'ar_comm','tp':'l','type':'main','label':''},
    {'f':'ar_comm', 'fp':'r','t':'ar_dest','tp':'l','type':'main','label':''},
    {'f':'ar_dest', 'fp':'r','t':'ar_uk',  'tp':'l','type':'yes', 'label':'UK/ROI'},
    {'f':'ar_dest', 'fp':'b','t':'ar_ce',  'tp':'t','type':'no',  'label':'CE'},
    {'f':'cus_css',    'fp':'r','t':'cus_ims',    'tp':'l','type':'main','label':''},
    {'f':'cus_ims',    'fp':'r','t':'cus_data',   'tp':'l','type':'main','label':''},
    {'f':'cus_data',   'fp':'r','t':'cus_sfd',    'tp':'l','type':'yes', 'label':'YES'},
    {'f':'cus_data',   'fp':'b','t':'cus_hold',   'tp':'t','type':'no',  'label':'NO'},
    {'f':'cus_hold',   'fp':'l','t':'cus_data',   'tp':'l','type':'loop','label':'resolved'},
    {'f':'cus_sfd',    'fp':'r','t':'cus_nbs',    'tp':'l','type':'main','label':''},
    {'f':'cus_nbs',    'fp':'r','t':'cus_accrual','tp':'l','type':'main','label':''},
    {'f':'cus_accrual','fp':'r','t':'cus_dda',    'tp':'l','type':'main','label':''},
]

# ── Slide groups (2 phases per slide) ────────────────────────────────────────
SLIDE_GROUPS = [
    {'phases':(0,1),'title':'Phase 1-2  |  Supplier Onboarding  →  Purchase Order Creation'},
    {'phases':(2,3),'title':'Phase 3-4  |  Shipment Execution  →  Invoice Preparation'},
    {'phases':(4,5),'title':'Phase 5-6  |  Cost Validation  →  PO Raise & Submission'},
    {'phases':(6,7),'title':'Phase 7-8  |  Payment & AR Billing  →  Customs Clearance & Duty'},
]

# ── Pre-compute element _cy ───────────────────────────────────────────────────
_LANE_MAP = {l['id']: l for l in LANES}
_ELEM_MAP = {}

def resolve():
    for e in ELEMS:
        lane = _LANE_MAP[e['lane']]
        cy   = (lane['y_top'] + lane['y_bot']) / 2 + e.get('cy_off', 0)
        e['_cy']   = cy
        e['_lane'] = lane
        _ELEM_MAP[e['id']] = e

resolve()

def get_port(elem, p):
    cx, cy, t = elem['cx'], elem['_cy'], elem['type']
    if t == 'dia':
        return {'r':(cx+DW,cy),'l':(cx-DW,cy),'t':(cx,cy-DW),'b':(cx,cy+DW)}[p]
    w = EW if t=='exc' else BW
    h = EH if t=='exc' else BH
    return {'r':(cx+w/2,cy),'l':(cx-w/2,cy),'t':(cx,cy-h/2),'b':(cx,cy+h/2)}[p]

# ── Coordinate transform ──────────────────────────────────────────────────────
SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)
TITLE_H = Inches(0.52)   # top title bar height

def make_transform(sg):
    pi, pj   = sg['phases']
    x_start  = PHASES[pi]['x_start']
    x_end    = PHASES[pj]['x_end']
    vis_w    = LABEL_W + (x_end - x_start)   # diagram pixels visible
    scale    = SLIDE_W / vis_w                # EMU per px (x, uniform)
    avail_h  = SLIDE_H - TITLE_H
    scale_y  = avail_h / CANVAS_H
    return {'x_start': x_start, 'x_end': x_end,
            'scale': scale, 'scale_y': scale_y}

def dx(diagram_x, T):
    if diagram_x <= LABEL_W:
        return int(diagram_x * T['scale'])
    adj = diagram_x - T['x_start'] + LABEL_W
    return int(adj * T['scale'])

def dy(diagram_y, T):
    return int(TITLE_H + diagram_y * T['scale_y'])

def in_range(elem, T):
    return T['x_start'] <= elem['cx'] <= T['x_end']

# ── XML helpers ───────────────────────────────────────────────────────────────
def set_fill(shape, fill_rgb, alpha_pct=100):
    """Set solid fill with optional transparency (alpha_pct 0-100)."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    if alpha_pct < 100:
        # Modify XML to add transparency
        spPr = shape._element.spPr
        solidFill = spPr.find('.//' + qn('a:solidFill'))
        if solidFill is not None:
            srgbClr = solidFill.find(qn('a:srgbClr'))
            if srgbClr is not None:
                alpha_el = etree.SubElement(srgbClr, qn('a:alpha'))
                alpha_el.set('val', str(int(alpha_pct * 1000)))

def set_line(shape, line_rgb, width_pt=1.0, dash=None):
    shape.line.color.rgb = line_rgb
    shape.line.width = Pt(width_pt)
    if dash:
        # set dash style via XML
        ln = shape.line._get_or_add_ln()
        prstDash = etree.SubElement(ln, qn('a:prstDash'))
        prstDash.set('val', dash)

def set_no_line(shape):
    shape.line.fill.background()

def add_arrow_to_connector(conn_shape, color_rgb, width_pt=1.2):
    """Set connector line color and add arrowhead at end."""
    conn_shape.line.color.rgb = color_rgb
    conn_shape.line.width = Pt(width_pt)
    # Add arrow via lxml on the cxnSp element
    spPr = conn_shape._element.spPr
    ln = spPr.find(qn('a:ln'))
    if ln is None:
        ln = etree.SubElement(spPr, qn('a:ln'))
    # Arrowhead at the END of the connector (tailEnd)
    tailEnd = ln.find(qn('a:tailEnd'))
    if tailEnd is None:
        tailEnd = etree.SubElement(ln, qn('a:tailEnd'))
    tailEnd.set('type', 'arrow')
    tailEnd.set('w', 'med')
    tailEnd.set('len', 'med')
    headEnd = ln.find(qn('a:headEnd'))
    if headEnd is None:
        headEnd = etree.SubElement(ln, qn('a:headEnd'))
    headEnd.set('type', 'none')

def set_text(shape, lines, font_pt=7, bold=False,
             color=None, align=PP_ALIGN.CENTER, valign_middle=True):
    if color is None:
        color = C_WHITE
    tf = shape.text_frame
    tf.word_wrap = False
    if valign_middle:
        tf.vertical_anchor = 3   # MSO_ANCHOR.MIDDLE = 3

    for i, line in enumerate(lines):
        p = tf.paragraphs[i] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size  = Pt(font_pt)
        run.font.bold  = bold
        run.font.color.rgb = color

def add_connector(slide, x1, y1, x2, y2, conn_type, label=''):
    """Add elbow connector with arrowhead."""
    from pptx.enum.shapes import MSO_CONNECTOR_TYPE as CT
    ctype = CT.ELBOW if conn_type in ('hand','loop','no') else CT.STRAIGHT
    conn = slide.shapes.add_connector(ctype, x1, y1, x2, y2)
    add_arrow_to_connector(conn, CONN_COLORS[conn_type], width_pt=1.0 if conn_type=='main' else 1.2)
    if label:
        # Small floating label near midpoint
        mx = (x1+x2)//2
        my = (y1+y2)//2
        tb = slide.shapes.add_textbox(mx-Inches(0.35), my-Inches(0.12),
                                       Inches(0.7), Inches(0.18))
        tf = tb.text_frame
        p  = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.size = Pt(5.5)
        r.font.color.rgb = CONN_COLORS[conn_type]
    return conn

# ── Shape drawing ─────────────────────────────────────────────────────────────
def draw_elem(slide, elem, T):
    t = elem['type']
    cx_ = dx(elem['cx'], T)
    cy_ = dy(elem['_cy'], T)
    lane_hex = elem['_lane']['color']
    lane_rgb = rgb(lane_hex)
    lines = elem['label'].split('\n')
    bold  = elem.get('bold', False)

    # font size scales with x-scale
    fpt = max(5.5, round(T['scale'] / Inches(1) * 72 * 0.92, 1))
    fpt = min(fpt, 8.0)

    if t == 'dia':
        w = int(DW * 2 * T['scale'])
        h = int(DW * 2 * T['scale_y'])
        left = cx_ - w//2
        top  = cy_ - h//2
        shape = slide.shapes.add_shape(4, left, top, w, h)  # 4 = diamond
        set_fill(shape, C_PANEL, alpha_pct=95)
        set_line(shape, C_AMBER, 1.2)
        set_text(shape, lines, fpt, bold=False, color=C_WHITE)

    elif t == 'exc':
        w = int(EW * T['scale'])
        h = int(EH * T['scale_y'])
        left = cx_ - w//2
        top  = cy_ - h//2
        shape = slide.shapes.add_shape(5, left, top, w, h)  # 5 = rounded rect
        set_fill(shape, C_PANEL, alpha_pct=95)
        set_line(shape, C_ORANGE, 1.2)
        set_text(shape, lines, max(fpt-0.5, 5), bold=False, color=C_ORANGE)

    else:  # box / out
        w = int(BW * T['scale'])
        h = int(BH * T['scale_y'])
        left = cx_ - w//2
        top  = cy_ - h//2
        shape = slide.shapes.add_shape(5, left, top, w, h)  # rounded rect
        if t == 'out':
            set_fill(shape, C_PANEL, alpha_pct=95)
            set_line(shape, C_TEAL, 1.4)
        else:
            set_fill(shape, C_PANEL, alpha_pct=95)
            set_line(shape, lane_rgb, 1.0)
        set_text(shape, lines, fpt, bold=bold, color=C_WHITE)

    # Sys tag below the shape
    sys_txt = elem.get('sys','')
    if sys_txt:
        tb = slide.shapes.add_textbox(left, top + h + Pt(1),
                                       w, Pt(9))
        tf = tb.text_frame
        p  = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = sys_txt
        r.font.size = Pt(max(fpt-1.5, 4.5))
        r.font.color.rgb = C_GREY

# ── Lane backgrounds & labels ─────────────────────────────────────────────────
def draw_lanes(slide, T):
    label_w = int(LABEL_W * T['scale'])
    full_w  = SLIDE_W
    for lane in LANES:
        lr = rgb(lane['color'])
        y_top = dy(lane['y_top'], T)
        h     = dy(lane['y_bot'], T) - y_top

        # Full-width subtle lane band
        band = slide.shapes.add_shape(1, 0, y_top, full_w, h)
        set_fill(band, lr, alpha_pct=8)
        set_no_line(band)

        # Left colour accent bar
        bar = slide.shapes.add_shape(1, 0, y_top, Pt(3), h)
        bar.fill.solid(); bar.fill.fore_color.rgb = lr
        set_no_line(bar)

        # Lane label panel background
        panel = slide.shapes.add_shape(1, 0, y_top, label_w, h)
        set_fill(panel, C_BG, alpha_pct=90)
        set_no_line(panel)

        # Lane name text
        name_lines = lane['label'].split('\n')
        sys_lines  = lane.get('systems', [])
        name_h = int(h * 0.45)
        tb = slide.shapes.add_textbox(Pt(6), y_top + (h - name_h)//2,
                                       label_w - Pt(8), name_h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = 3
        for i, line in enumerate(name_lines):
            p = tf.paragraphs[i] if i==0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = line
            r.font.size = Pt(max(6, round(T['scale']/Inches(1)*72*0.7, 1)))
            r.font.bold = True
            r.font.color.rgb = lr

        # System tags (small, below name)
        sys_fpt = Pt(max(4.5, round(T['scale']/Inches(1)*72*0.52, 1)))
        for si, sys_name in enumerate(sys_lines):
            sy = y_top + int(h * 0.55) + si * int(sys_fpt * 1.6)
            if sy + sys_fpt > y_top + h - Pt(2):
                break
            stb = slide.shapes.add_textbox(Pt(6), sy, label_w - Pt(10), sys_fpt + Pt(2))
            stf = stb.text_frame
            sp  = stf.paragraphs[0]
            sp.alignment = PP_ALIGN.CENTER
            sr  = sp.add_run()
            sr.text = sys_name
            sr.font.size = sys_fpt
            sr.font.color.rgb = rgb('#4E6280')

        # Lane separator line
        ln_shape = slide.shapes.add_shape(1, 0, y_top + h - Pt(0.5), full_w, Pt(0.5))
        set_fill(ln_shape, rgb('#1E293B'), alpha_pct=100)
        set_no_line(ln_shape)

def draw_phase_headers(slide, sg, T):
    pi, pj = sg['phases']
    label_w = int(LABEL_W * T['scale'])
    for ph_i in range(pi, pj+1):
        ph = PHASES[ph_i]
        x1 = dx(ph['x_start'], T)
        x2 = dx(ph['x_end'],   T)
        pw = x2 - x1
        # Phase background top strip
        hdr = slide.shapes.add_shape(1, x1, 0, pw, TITLE_H)
        phase_colors = ['#00E5C0','#4FB3FF','#A78BFA','#FFD04A',
                        '#FF6B6B','#C57BFF','#52E67A','#38C6FF']
        set_fill(hdr, rgb(phase_colors[ph_i]), alpha_pct=12)
        set_no_line(hdr)
        # Phase label
        lines = ph['label'].split('\n')
        tb = slide.shapes.add_textbox(x1 + Pt(4), 0, pw - Pt(8), TITLE_H)
        tf = tb.text_frame
        tf.vertical_anchor = 3
        for i, line in enumerate(lines):
            p = tf.paragraphs[i] if i==0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = line
            r.font.size  = Pt(7)
            r.font.bold  = True
            r.font.color.rgb = rgb(phase_colors[ph_i])
        # Vertical separator
        sep = slide.shapes.add_shape(1, x2-Pt(0.5), 0, Pt(0.5), SLIDE_H)
        set_fill(sep, rgb('#1E293B'), alpha_pct=100)
        set_no_line(sep)

    # Label panel header
    lb = slide.shapes.add_textbox(0, 0, int(LABEL_W * T['scale']), TITLE_H)
    tf = lb.text_frame
    tf.vertical_anchor = 3
    p  = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r  = p.add_run()
    r.text = 'SWIM LANE  /  SYSTEM'
    r.font.size = Pt(6)
    r.font.bold = True
    r.font.color.rgb = C_GREY

def draw_slide_title(slide, title_text):
    """Slide number + title in top-left."""
    tb = slide.shapes.add_textbox(Pt(6), Pt(4), Inches(10), Pt(18))
    tf = tb.text_frame
    p  = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r  = p.add_run()
    r.text = title_text
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = C_TEAL

def draw_background(slide):
    bg = slide.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = C_BG
    set_no_line(bg)
    sp = bg._element
    sp.getparent().remove(sp)
    slide.shapes._spTree.insert(2, sp)

# ── Connections ───────────────────────────────────────────────────────────────
def draw_connections(slide, T):
    for conn in CONNS:
        fe = _ELEM_MAP.get(conn['f'])
        te = _ELEM_MAP.get(conn['t'])
        if fe is None or te is None:
            continue
        # Only draw if at least FROM element is in range; clamp TO endpoint
        if not in_range(fe, T):
            continue
        fp = get_port(fe, conn['fp'])
        # If target not in range, clamp to slide edge
        if in_range(te, T):
            tp = get_port(te, conn['tp'])
        else:
            tp_raw = get_port(te, conn['tp'])
            # clamp x to slide content edge
            max_diag_x = T['x_end']
            clamp_x_diag = min(tp_raw[0], max_diag_x)
            tp = (clamp_x_diag, tp_raw[1])

        x1 = dx(fp[0], T); y1 = dy(fp[1], T)
        x2 = dx(tp[0], T); y2 = dy(tp[1], T)

        # Skip zero-length connectors
        if abs(x2-x1) < 3 and abs(y2-y1) < 3:
            continue

        add_connector(slide, x1, y1, x2, y2,
                      conn['type'], label=conn.get('label',''))

# ── Build one diagram slide ───────────────────────────────────────────────────
def build_diagram_slide(prs, sg, idx):
    blank = prs.slide_layouts[6]
    sl    = prs.slides.add_slide(blank)
    T     = make_transform(sg)

    draw_background(sl)
    draw_lanes(sl, T)
    draw_phase_headers(sl, sg, T)
    draw_slide_title(sl, f"Slide {idx+2}  |  {sg['title']}")
    draw_connections(sl, T)   # draw connections BEFORE shapes so arrows go behind

    for elem in ELEMS:
        if in_range(elem, T):
            draw_elem(sl, elem, T)

    print("  Slide " + str(idx+2) + " built")

# ── Title slide ───────────────────────────────────────────────────────────────
def build_title_slide(prs):
    blank = prs.slide_layouts[6]
    sl    = prs.slides.add_slide(blank)

    draw_background(sl)

    # Teal accent bar
    bar = sl.shapes.add_shape(1, 0, Inches(3.1), SLIDE_W, Pt(1.5))
    bar.fill.solid(); bar.fill.fore_color.rgb = C_TEAL; set_no_line(bar)

    def txt(t, x, y, w, h, size, bold=False, color=C_WHITE, align=PP_ALIGN.CENTER):
        tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True
        tf.vertical_anchor = 3
        p  = tf.paragraphs[0]; p.alignment = align
        r  = p.add_run(); r.text = t
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color

    txt('GLO Finance', 1, 1.0, 11, 0.8, 42, bold=True, color=C_TEAL)
    txt('AS-IS Process Map', 1, 1.85, 11, 0.65, 28, color=C_WHITE)
    txt('7 Swim Lanes  ·  8 Process Phases  ·  Document Source: GLO_Finance_ASIS_Final.docx',
        1, 2.55, 11, 0.4, 11, color=C_GREY)
    txt('Tesco Global Finance  |  Confidential  |  2024',
        0, 7.15, 13.33, 0.3, 8, color=C_GREY)

    # Legend
    legend = [
        ('Process Step', C_GREY,   'box'),
        ('Key Output',   C_TEAL,   'out'),
        ('Decision',     C_AMBER,  'dia'),
        ('Exception',    C_ORANGE, 'exc'),
        ('YES path',     C_GREEN,  'line'),
        ('NO / Exception',C_ORANGE,'line'),
        ('Cross-lane',   rgb('#64748B'), 'line'),
    ]
    lx = Inches(1.3)
    for i, (lbl, clr, shape_type) in enumerate(legend):
        lx_i = lx + i * Inches(1.7)
        if shape_type in ('box','out','exc'):
            sw = Pt(28); sh = Pt(14)
            s = sl.shapes.add_shape(5, lx_i, Inches(3.55), sw, sh)
            s.fill.solid(); s.fill.fore_color.rgb = C_PANEL
            s.line.color.rgb = clr; s.line.width = Pt(1.2)
        elif shape_type == 'dia':
            sw = Pt(16); sh = Pt(16)
            s = sl.shapes.add_shape(4, lx_i, Inches(3.55), sw, sh)
            s.fill.solid(); s.fill.fore_color.rgb = C_PANEL
            s.line.color.rgb = clr; s.line.width = Pt(1.2)
        else:
            s = sl.shapes.add_shape(1, lx_i, Inches(3.62), Pt(28), Pt(2))
            s.fill.solid(); s.fill.fore_color.rgb = clr; set_no_line(s)
        tb = sl.shapes.add_textbox(lx_i - Pt(2), Inches(3.82), Pt(90), Pt(12))
        tf = tb.text_frame; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = lbl; r.font.size = Pt(8); r.font.color.rgb = C_GREY

    # Lane summary table
    ty = 4.3
    for lane in LANES:
        lr = rgb(lane['color'])
        dot = sl.shapes.add_shape(9, Inches(1.3), Inches(ty)+Pt(2), Pt(10), Pt(10))
        dot.fill.solid(); dot.fill.fore_color.rgb = lr; set_no_line(dot)
        tb = sl.shapes.add_textbox(Inches(1.55), Inches(ty), Inches(5.5), Pt(14))
        tf = tb.text_frame; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = lane['label'].replace('\n',' ') + '   –   ' + ' · '.join(lane['systems'])
        r.font.size = Pt(8); r.font.color.rgb = C_WHITE
        ty += 0.43

    print("  Title slide built")

# ── Main ──────────────────────────────────────────────────────────────────────
def main():
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H

    build_title_slide(prs)
    for i, sg in enumerate(SLIDE_GROUPS):
        build_diagram_slide(prs, sg, i)

    prs.save(OUTPUT)
    print("Saved: " + str(OUTPUT))

if __name__ == '__main__':
    main()
